#!/usr/bin/env python3
"""
Gera um conjunto de slides explicando o trabalho HyMRPL, alinhado com a
dissertação escrita. Ênfase no motor adaptativo, na normalização dos dados
e na janela deslizante.

Saída: HyMRPL_Visao_Geral.pptx
Requer: python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# --- Paleta ---
NAVY   = RGBColor(0x0B, 0x29, 0x47)
BLUE   = RGBColor(0x1F, 0x6F, 0xB2)
TEAL   = RGBColor(0x12, 0x8C, 0x7D)
LIGHT  = RGBColor(0xED, 0xF2, 0xF7)
GREY   = RGBColor(0x44, 0x4A, 0x52)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xE8, 0x7A, 0x1E)

W, H = Inches(13.333), Inches(7.5)


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # em branco


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def box(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h).text_frame


def style(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, italic=False):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return p


def bullets(tf, items, size=18, color=GREY, gap=6):
    tf.word_wrap = True
    first = True
    for it in items:
        lvl = 0
        txt = it
        if isinstance(it, tuple):
            lvl, txt = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = ("• " if lvl == 0 else "– ") + txt
        r.font.size = Pt(size - lvl * 2)
        r.font.color.rgb = color
        r.font.name = "Calibri"


def bar(slide, color=BLUE, top=Inches(1.18), h=Inches(0.06)):
    s = slide.shapes.add_shape(1, Inches(0.6), top, Inches(12.1), h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()


def header(slide, title, kicker=None):
    bg(slide, WHITE)
    if kicker:
        style(box(slide, Inches(0.6), Inches(0.28), Inches(12), Inches(0.4)),
              kicker, 13, ACCENT, bold=True)
    style(box(slide, Inches(0.6), Inches(0.55), Inches(12.1), Inches(0.7)),
          title, 30, NAVY, bold=True)
    bar(slide)


def card(slide, l, t, w, h, fill, title, lines, tcolor=WHITE, bcolor=None):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(12); tf.margin_right = Pt(12)
    tf.margin_top = Pt(10); tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.bold = True; r.font.size = Pt(16); r.font.color.rgb = tcolor
    r.font.name = "Calibri"
    for ln in lines:
        pp = tf.add_paragraph()
        pp.space_before = Pt(4)
        rr = pp.add_run(); rr.text = ln
        rr.font.size = Pt(12.5)
        rr.font.color.rgb = bcolor if bcolor else tcolor
        rr.font.name = "Calibri"
    return s


prs = Presentation()
prs.slide_width = W
prs.slide_height = H

# ============================================================
# 1. TÍTULO
# ============================================================
s = add_slide(prs); bg(s, NAVY)
band = s.shapes.add_shape(1, 0, Inches(2.5), W, Inches(2.05))
band.fill.solid(); band.fill.fore_color.rgb = BLUE; band.line.fill.background()
style(box(s, Inches(0.8), Inches(2.62), Inches(11.7), Inches(1.0)),
      "HyMRPL — Modo Híbrido de Operação para o RPL (MOP = 6)", 34, WHITE, bold=True)
style(box(s, Inches(0.8), Inches(3.6), Inches(11.7), Inches(0.7)),
      "Coexistência simultânea de comportamentos Storing e Non-Storing em uma mesma DODAG",
      18, LIGHT)
style(box(s, Inches(0.8), Inches(5.0), Inches(11.7), Inches(0.5)),
      "Adaptação de classe em tempo de execução, sem overhead na rede  ·  implementado no daemon rpld",
      15, RGBColor(0xBF,0xD6,0xEA))
style(box(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.5)),
      "Cassius Clay Batista da Silva Filho  ·  PPgTI", 14, RGBColor(0x9F,0xB6,0xCC))

# ============================================================
# 2. PROBLEMA E MOTIVAÇÃO
# ============================================================
s = add_slide(prs)
header(s, "O Problema: o MOP Global Impõe Homogeneidade", "Motivação")
tf = box(s, Inches(0.7), Inches(1.5), Inches(6.0), Inches(5.4))
bullets(tf, [
    "O RPL (RFC 6550) define o Modo de Operação por DODAG — todos os nós usam o mesmo modo.",
    "Storing: rotas descendentes em cada nó — baixa latência, mas custo de memória nos nós restritos.",
    "Non-Storing: estado centralizado na raiz via SRH — nós leves, porém caminhos mais longos.",
    "Redes IoT reais são heterogêneas: energia, memória e mobilidade variam de nó para nó.",
    "Um único modo global não acomoda essa assimetria.",
], size=17)
card(s, Inches(7.0), Inches(1.6), Inches(5.6), Inches(4.6), LIGHT,
     "Ideia do HyMRPL", [
        "MOP = 6 (experimental, a RFC reserva 4–7).",
        "Cada nó escolhe localmente sua classe funcional:",
        "   Classe S → storing-like (mantém rotas locais)",
        "   Classe N → non-storing-like (delega ao SRH da raiz)",
        "Os dois paradigmas coexistem na mesma DODAG.",
        "Troca de classe em tempo de execução, sem pacotes extras.",
     ], tcolor=NAVY, bcolor=GREY)

# ============================================================
# 3. ARQUITETURA
# ============================================================
s = add_slide(prs)
header(s, "Arquitetura: Três Subsistemas Integrados", "Visão geral")
card(s, Inches(0.7), Inches(1.6), Inches(3.85), Inches(4.7), BLUE,
     "1. Roteamento Híbrido (MOP=6)", [
        "DAO dual: para o parent e para a raiz.",
        "A raiz constrói a árvore SRH completa.",
        "Classe S instala rotas hop-by-hop.",
        "Classe N delega ao SRH da raiz.",
        "Classe S agrega os alvos dos filhos",
        "→ enriquece a árvore SRH da raiz.",
     ])
card(s, Inches(4.75), Inches(1.6), Inches(3.85), Inches(4.7), TEAL,
     "2. Motor de Decisão Adaptativa", [
        "Roda dentro do event loop do rpld.",
        "Score composto de 3 métricas:",
        "PDR, energia residual, estabilidade.",
        "Score normalizado + histerese.",
        "Sem processo Python externo.",
        "Mudança de 1 byte → zero overhead.",
     ])
card(s, Inches(8.8), Inches(1.6), Inches(3.85), Inches(4.7), NAVY,
     "3. FIFO Seguro", [
        "Autenticação HMAC-SHA256.",
        "Nonce → proteção contra replay.",
        "Rate limiting (10s / 3 por min).",
        "Permissões 0600 (somente root).",
        "Log de auditoria de cada tentativa.",
        "Fallback legado em texto plano.",
     ])
style(box(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.5)),
      "Comandos externos via FIFO têm prioridade sobre o motor adaptativo (override manual por operadores / SDN).",
      14, GREY, italic=True)

# ============================================================
# 4. MOTOR ADAPTATIVO — SCORE COMPOSTO
# ============================================================
s = add_slide(prs)
header(s, "Motor de Decisão Adaptativa: o Score Composto", "Modelo central")
fb = s.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(11.9), Inches(1.15))
fb.fill.solid(); fb.fill.fore_color.rgb = LIGHT
fb.line.fill.background()
tf = fb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
style(tf, "S  =  0,4 · s_pdr   +   0,3 · s_energia   +   0,3 · s_estab        →    S ≥ 0,75 ⇒ Classe S    |    S < 0,75 ⇒ Classe N",
      18, NAVY, bold=True, align=PP_ALIGN.CENTER)
tf2 = box(s, Inches(0.7), Inches(2.95), Inches(11.9), Inches(0.5))
bullets(tf2, ["Cada componente é normalizado em [0, 1] e os pesos somam 1,0 → o próprio score fica em [0, 1] e é diretamente comparável ao limiar θ = 0,75."], size=15, color=GREY)

card(s, Inches(0.7), Inches(3.7), Inches(3.85), Inches(2.9), WHITE,
     "s_pdr  (peso 0,4)", [
        "Qualidade do enlace.",
        "Fração de DAO-ACKs recebidos",
        "em janela deslizante de W = 20.",
        "Reage em escala de segundos.",
     ], tcolor=BLUE, bcolor=GREY)
card(s, Inches(4.75), Inches(3.7), Inches(3.85), Inches(2.9), WHITE,
     "s_energia  (peso 0,3)", [
        "Bateria / capacidade residual.",
        "Lida a cada 5 s, suavizada por",
        "EMA com α = 0,3.",
        "Varia lentamente (min–horas).",
     ], tcolor=TEAL, bcolor=GREY)
card(s, Inches(8.8), Inches(3.7), Inches(3.85), Inches(2.9), WHITE,
     "s_estab  (peso 0,3)", [
        "Estabilidade topológica.",
        "Frequência de troca de parent em",
        "janela de 60 s, normalizada por Nmax=3.",
        "Índice contínuo, não binário.",
     ], tcolor=NAVY, bcolor=GREY)
for x, c in [(Inches(0.7), BLUE), (Inches(4.75), TEAL), (Inches(8.8), NAVY)]:
    tab = s.shapes.add_shape(1, x, Inches(3.7), Inches(3.85), Inches(0.12))
    tab.fill.solid(); tab.fill.fore_color.rgb = c; tab.line.fill.background()

# ============================================================
# 5. NORMALIZAÇÃO DOS DADOS (CHAVE)
# ============================================================
s = add_slide(prs)
header(s, "Normalização dos Dados: Comparando Grandezas Distintas", "Por que importa")
tf = box(s, Inches(0.7), Inches(1.5), Inches(6.1), Inches(5.5))
bullets(tf, [
    "As três métricas têm naturezas e unidades completamente diferentes:",
    (1, "PDR → razão de sucesso do protocolo (%)"),
    (1, "Energia → carga da bateria (%)"),
    (1, "Estabilidade → contagem de eventos topológicos"),
    "Somá-las diretamente não faria sentido — uma poderia dominar só pela sua escala.",
    "Cada métrica é mapeada para a escala comum [0, 1]:",
    (1, "s_pdr = sucessos / W"),
    (1, "s_energia = EMA( E(t)/100 )"),
    (1, "s_estab = 1 − min(1, N_trocas / Nmax)"),
    "Como todo s_k ∈ [0,1] e Σ w_k = 1, o score S também permanece em [0,1].",
], size=15)
card(s, Inches(7.0), Inches(1.6), Inches(5.6), Inches(2.45), LIGHT,
     "O que a normalização garante", [
        "Contribuições proporcionais e comparáveis.",
        "Um único limiar que preserva a escala (θ=0,75).",
        "Os pesos expressam prioridade, não magnitude.",
        "Robustez a mudanças de unidade/faixa do sensor.",
     ], tcolor=NAVY, bcolor=GREY)
card(s, Inches(7.0), Inches(4.25), Inches(5.6), Inches(2.35), NAVY,
     "Exemplo (só a bateria varia)", [
        "Nó estável, bateria cheia:",
        "S = 0,4·1 + 0,3·1,0 + 0,3·1 = 1,00 → S",
        "Mesmo nó, bateria 10%:",
        "S = 0,4·1 + 0,3·0,10 + 0,3·1 = 0,73 → N",
     ])

# ============================================================
# 6. JANELA DESLIZANTE (CHAVE)
# ============================================================
s = add_slide(prs)
header(s, "Janela Deslizante: Estimativa de PDR Recente e Estável", "Por que importa")
tf = box(s, Inches(0.7), Inches(1.5), Inches(6.1), Inches(5.5))
bullets(tf, [
    "O PDR é estimado a partir dos últimos W = 20 resultados de DAO-ACK (sucesso = 1, timeout = 0).",
    "s_pdr = (1/W) · Σ 1[DAO-ACK_k recebido]",
    "Implementada como buffer circular: a amostra mais nova sobrescreve a mais antiga (índice em módulo W).",
    "É uma média móvel sobre um histórico recente de tamanho fixo.",
    "Por que janela e não média desde o início:",
    (1, "A média acumulada reage cada vez mais devagar — amostras antigas nunca saem."),
    (1, "A janela mantém a estimativa atual e filtra perdas pontuais."),
    (1, "W = 20 ≈ 20 s de histórico: estável o bastante e ainda responsivo."),
], size=15)
style(box(s, Inches(7.0), Inches(1.5), Inches(5.6), Inches(0.4)),
      "Buffer circular (W = 20 amostras)", 14, NAVY, bold=True)
cell_w = Inches(0.52)
for i in range(10):
    cx = Emu(int(Inches(7.0)) + i * int(cell_w))
    c = s.shapes.add_shape(1, cx, Inches(2.0), cell_w, Inches(0.52))
    filled = i not in (3,)
    c.fill.solid(); c.fill.fore_color.rgb = TEAL if filled else ACCENT
    c.line.color.rgb = WHITE
    cc = c.text_frame; cc.word_wrap = False
    pp = cc.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = "1" if filled else "0"
    rr.font.size = Pt(12); rr.font.bold = True; rr.font.color.rgb = WHITE
for i in range(10):
    cx = Emu(int(Inches(7.0)) + i * int(cell_w))
    c = s.shapes.add_shape(1, cx, Inches(2.54), cell_w, Inches(0.52))
    filled = i not in (6,)
    c.fill.solid(); c.fill.fore_color.rgb = TEAL if filled else ACCENT
    c.line.color.rgb = WHITE
    cc = c.text_frame
    pp = cc.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = "1" if filled else "0"
    rr.font.size = Pt(12); rr.font.bold = True; rr.font.color.rgb = WHITE
style(box(s, Inches(7.0), Inches(3.15), Inches(5.6), Inches(0.4)),
      "18 sucessos / 20  →  s_pdr = 0,90", 15, NAVY, bold=True)
card(s, Inches(7.0), Inches(3.75), Inches(5.6), Inches(2.85), LIGHT,
     "Janela vs. EMA — ferramentas distintas", [
        "PDR: janela deslizante aritmética (peso igual,",
        "  horizonte recente fixo, rápida porém estável).",
        "Energia: EMA α=0,3 (desvanecimento exponencial,",
        "  filtra ruído do sensor em sinal lento).",
        "Estabilidade: contagem normalizada em 60 s.",
        "Cada janela combina com a dinâmica da métrica.",
     ], tcolor=NAVY, bcolor=GREY)

# ============================================================
# 7. HISTERESE + FIFO SEGURO
# ============================================================
s = add_slide(prs)
header(s, "Do Score à Ação: Histerese e Controle Seguro", "Fluxo de decisão")
card(s, Inches(0.7), Inches(1.6), Inches(5.85), Inches(4.9), LIGHT,
     "Histerese (anti-oscilação)", [
        "A troca só é aplicada após 3 ciclos consecutivos",
        "recomendando a mesma nova classe.",
        "Evita flapping quando o score fica perto de θ.",
        "O contador zera se a recomendação inverte.",
        "Mais o rate limiting compartilhado: ≥ 10 s entre",
        "trocas, ≤ 3 por minuto.",
        "",
        "Resultado: trocas deliberadas, nunca ruidosas.",
     ], tcolor=NAVY, bcolor=GREY)
card(s, Inches(6.7), Inches(1.6), Inches(5.9), Inches(4.9), NAVY,
     "FIFO Seguro — 5 camadas", [
        "1. HMAC-SHA256 com token de 256 bits.",
        "2. Nonce monotônico → bloqueia replay.",
        "3. Rate limiting (compartilhado com o motor).",
        "4. FIFO modo 0600 → só root pode escrever.",
        "5. Log de auditoria (aceitos/rejeitados).",
        "",
        "CLASS_S | <nonce_hex> | <hmac_hex>",
        "Fallback em texto plano quando não há token.",
     ])

# ============================================================
# 8. VALIDAÇÃO EXPERIMENTAL
# ============================================================
s = add_slide(prs)
header(s, "Validação Experimental (Mininet-WiFi + 6LoWPAN)", "Resultados")
card(s, Inches(0.7), Inches(1.6), Inches(5.85), Inches(2.3), LIGHT,
     "Decisão adaptativa (6 fases)", [
        "Classe correta em todas as fases:",
        "Estável + bateria cheia → S (score ≈ 1,00).",
        "Bateria baixa → N (score ≈ 0,73).",
        "Recuperação → volta para S.",
     ], tcolor=NAVY, bcolor=GREY)
card(s, Inches(6.7), Inches(1.6), Inches(5.9), Inches(2.3), LIGHT,
     "Troca dinâmica via FIFO seguro", [
        "N→S→N em execução, 100% de PDR o tempo todo.",
        "Latência local s4→s5 −38,6% após N→S.",
        "Comandos autenticados, sem pacotes extras.",
     ], tcolor=NAVY, bcolor=GREY)
card(s, Inches(0.7), Inches(4.05), Inches(5.85), Inches(2.4), TEAL,
     "Resiliência em mesh (falha após restauração)", [
        "PDR:  HyMRPL 87,2%  vs  Storing 79,5%",
        "        vs  Non-Storing 64,1%.",
        "+23,1 pp sobre Non-Storing, +7,7 sobre Storing.",
        "Árvore SRH enriquecida → mais caminhos alternativos.",
     ])
card(s, Inches(6.7), Inches(4.05), Inches(5.9), Inches(2.4), BLUE,
     "Custo", [
        "CPU, memória e tráfego de controle idênticos",
        "aos modos tradicionais.",
        "Troca de classe é puramente local (1 byte).",
        "Adaptação sem overhead mensurável.",
     ])

# ============================================================
# 9. VERIFICAÇÃO DO CÓDIGO + ENCERRAMENTO
# ============================================================
s = add_slide(prs)
header(s, "Verificação do Código e Conclusões", "Encerramento")
tf = box(s, Inches(0.7), Inches(1.5), Inches(7.1), Inches(5.3))
bullets(tf, [
    "A implementação em C reproduz fielmente o modelo da dissertação:",
    (1, "Janela deslizante W=20, EMA α=0,3, índice de estabilidade, pesos 0,4/0,3/0,3, θ=0,75, histerese de 3 ciclos."),
    "Os hooks de PDR (DAO-ACK) e de troca de parent foram integrados ao process.c — sem alterar mensagens RPL nem emitir pacotes.",
    "A detecção de timeout de DAO-ACK é local; o FIFO usa comparação HMAC em tempo constante e nonce crescente.",
    "hymrpl_cmd.c compila limpo com -Wall -Wextra; sem diagnósticos no motor, no header e no daemon.",
    "Experimentos em Python validados nos fluxos de bateria e de FIFO.",
], size=15)
card(s, Inches(8.0), Inches(1.6), Inches(4.6), Inches(4.9), NAVY,
     "Principais contribuições", [
        "Coexistência S/N por nó em uma DODAG (MOP=6).",
        "Motor adaptativo de score composto normalizado.",
        "Controle de classe seguro e autenticado em runtime.",
        "Ganho de resiliência sem overhead na rede.",
     ])

prs.save("HyMRPL_Visao_Geral.pptx")
print("Salvo HyMRPL_Visao_Geral.pptx")
